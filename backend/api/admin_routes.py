import csv
from datetime import date, datetime, timedelta
import io
import json
import os
import time
from collections import defaultdict
from pathlib import Path

from database.connection import get_db
from database.models import (
    Administrador,
    EstadoJustificacion,
    LogAuditoria,
    Marcaje,
    Sucursal,
    TipoAccion,
    TipoMarcaje,
    Usuario,
)
from fastapi import APIRouter, Depends, Form, HTTPException, Request, Response
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from jose import JWTError, jwt
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from passlib.context import CryptContext
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import func as sqlfunc
from sqlalchemy.orm import Session

router = APIRouter()

# ── Configuración de ruta absoluta para plantillas ───────────────────────
BASE_DIR = Path(__file__).resolve().parent.parent  # Apunta a /code/backend
templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))

# ══════════════════════════════════════════
# CONFIGURACIÓN AUTH
# ══════════════════════════════════════════

SECRET_KEY = os.getenv("SECRET_KEY", "cambia-esto-urgente")
ALGORITHM = "HS256"
ACCESS_TOKEN_MINUTES = int(os.getenv("ACCESS_TOKEN_EXPIRE_MINUTES", 480))
SUPERADMIN_EMAIL = os.getenv("SUPERADMIN_EMAIL", "soporte@tuempresa.com")
SUPERADMIN_PASSWORD = os.getenv("SUPERADMIN_PASSWORD", "cambia-esto")

pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")

# ── Rate limiting en memoria (login) ───────────────────────────────────
_login_intentos: dict[str, list] = defaultdict(list)
MAX_INTENTOS_LOGIN = 10
VENTANA_LOGIN_SEG = 300  # 5 minutos


def _ip(request: Request) -> str:
    """Extrae la IP real considerando proxies."""
    forwarded = request.headers.get("X-Forwarded-For")
    if forwarded:
        return forwarded.split(",")[0].strip()
    return request.client.host if request.client else "127.0.0.1"


def _check_rate_limit_login(ip: str):
    """Lanza 429 si la IP superó el límite de intentos de login."""
    ahora = time.time()
    recientes = [t for t in _login_intentos[ip] if ahora - t < VENTANA_LOGIN_SEG]
    _login_intentos[ip] = recientes
    if len(recientes) >= MAX_INTENTOS_LOGIN:
        raise HTTPException(
            status_code=429,
            detail=f"Demasiados intentos. Espera {VENTANA_LOGIN_SEG // 60} minutos.",
        )
    _login_intentos[ip].append(ahora)


# ══════════════════════════════════════════
# HELPERS AUTH
# ══════════════════════════════════════════

def verificar_password(plain: str, hashed: str) -> bool:
    return pwd_context.verify(plain, hashed)


def hashear_password(plain: str) -> str:
    return pwd_context.hash(plain)


def crear_token(data: dict, minutos: int = ACCESS_TOKEN_MINUTES) -> str:
    payload = data.copy()
    payload["exp"] = datetime.utcnow() + timedelta(minutes=minutos)
    payload["iat"] = datetime.utcnow()
    return jwt.encode(payload, SECRET_KEY, algorithm=ALGORITHM)


def obtener_admin_actual(request: Request, db: Session = Depends(get_db)) -> dict:
    """
    Lee el JWT de la cookie, lo valida y retorna el dict del admin
    con las sucursales asignadas.
    """
    token = request.cookies.get("biocore_token")
    if not token:
        raise HTTPException(status_code=401, detail="No autenticado")
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        email = payload.get("sub")
        es_super = payload.get("superadmin", False)
    except JWTError as e:
        raise HTTPException(status_code=401, detail=f"Sesión inválida o expirada: {e}")

    if es_super:
        sucursales_ids = [s.id for s in db.query(Sucursal).all()]
        return {
            "id": 0,
            "nombre": "Soporte BioCore",
            "email": email,
            "es_superadmin": True,
            "sucursales_ids": sucursales_ids,
        }

    admin = db.query(Administrador).filter(Administrador.email == email).first()
    if not admin:
        raise HTTPException(status_code=401, detail="Administrador no encontrado")

    sucursales_ids = [s.id for s in admin.sucursales] if hasattr(admin, "sucursales") else []

    return {
        "id": admin.id,
        "nombre": admin.nombre,
        "email": admin.email,
        "es_superadmin": getattr(admin, "es_superadmin", False),
        "sucursales_ids": sucursales_ids,
    }


def requiere_superadmin(admin=Depends(obtener_admin_actual)) -> dict:
    """Middleware/Dependencia para restringir rutas a sólo superadministradores."""
    if not admin.get("es_superadmin"):
        raise HTTPException(status_code=403, detail="Solo superadministradores")
    return admin


# ══════════════════════════════════════════
# HELPER AUDITORÍA (CORREGIDO PARA EVITAR ERROR 500)
# ══════════════════════════════════════════

def _log(
    db: Session,
    admin: dict,
    accion: TipoAccion,
    detalle: dict = None,
    ip: str = None,
):
    """Registra una acción de admin en logs_auditoria convirtiendo 'detalle' dict a JSON string."""
    detalle_str = json.dumps(detalle) if isinstance(detalle, dict) else detalle
    entrada = LogAuditoria(
        admin_id=admin["id"] if admin["id"] != 0 else None,
        admin_email=admin["email"],
        accion=accion,
        detalle=detalle_str,
        ip=ip,
    )
    db.add(entrada)
    db.commit()


# ══════════════════════════════════════════
# LOGIN / LOGOUT
# ══════════════════════════════════════════

@router.get("/admin/login", response_class=HTMLResponse)
@router.get("/admin_login", response_class=HTMLResponse)
async def login_page(request: Request):
    return templates.TemplateResponse(request, "admin_login.html")


@router.post("/admin/login")
async def login(
    request: Request,
    email: str = Form(...),
    password: str = Form(...),
    db: Session = Depends(get_db),
):
    ip = _ip(request)
    _check_rate_limit_login(ip)

    # Superadmin maestro (credenciales desde .env)
    if email == SUPERADMIN_EMAIL and password == SUPERADMIN_PASSWORD:
        token = crear_token({"sub": email, "superadmin": True})
        resp = RedirectResponse(url="/admin/dashboard", status_code=302)
        resp.set_cookie(
            "biocore_token",
            token,
            httponly=True,
            samesite="lax",
            max_age=ACCESS_TOKEN_MINUTES * 60,
        )
        _log(
            db,
            {"id": 0, "email": email, "es_superadmin": True},
            TipoAccion.LOGIN_ADMIN,
            {"via": "superadmin"},
            ip,
        )
        return resp

    # Admin normal
    admin = db.query(Administrador).filter(Administrador.email == email).first()
    if not admin or not verificar_password(password, admin.password_hash):
        return templates.TemplateResponse(
            request, "admin_login.html", {"error": "Credenciales incorrectas"}
        )

    # Actualizar último login
    if hasattr(admin, "ultimo_login"):
        admin.ultimo_login = datetime.utcnow()
        db.commit()

    es_superadmin_flag = getattr(admin, "es_superadmin", False)
    token = crear_token({"sub": admin.email, "superadmin": es_superadmin_flag})
    resp = RedirectResponse(url="/admin/dashboard", status_code=302)
    resp.set_cookie(
        "biocore_token",
        token,
        httponly=True,
        samesite="lax",
        max_age=ACCESS_TOKEN_MINUTES * 60,
    )
    _log(
        db,
        {"id": admin.id, "email": admin.email, "es_superadmin": es_superadmin_flag},
        TipoAccion.LOGIN_ADMIN,
        {"via": "admin_normal"},
        ip,
    )
    return resp


@router.get("/admin/logout")
async def logout(
    request: Request,
    db: Session = Depends(get_db),
    admin=Depends(obtener_admin_actual),
):
    _log(db, admin, TipoAccion.LOGOUT_ADMIN, ip=_ip(request))
    resp = RedirectResponse(url="/admin/login", status_code=302)
    resp.delete_cookie("biocore_token")
    return resp


# ══════════════════════════════════════════
# DASHBOARD
# ══════════════════════════════════════════

@router.get("/admin/dashboard", response_class=HTMLResponse)
async def dashboard(
    request: Request,
    db: Session = Depends(get_db),
    admin=Depends(obtener_admin_actual),
):
    hoy = date.today()
    sucursales_ids = admin.get("sucursales_ids", [])

    query_usuarios = db.query(Usuario).filter(Usuario.activo == True)
    query_marcajes = db.query(Marcaje).filter(sqlfunc.date(Marcaje.fecha_hora) == hoy)
    query_justif = db.query(Marcaje).filter(
        Marcaje.estado_justificacion == EstadoJustificacion.PENDIENTE,
        Marcaje.es_especial == True,
    )
    query_entradas = db.query(Marcaje).filter(
        sqlfunc.date(Marcaje.fecha_hora) == hoy,
        Marcaje.tipo == TipoMarcaje.ENTRADA,
    )
    query_tardanzas = db.query(sqlfunc.count(Marcaje.id)).filter(
        sqlfunc.date(Marcaje.fecha_hora) == hoy,
        Marcaje.tipo == TipoMarcaje.ENTRADA,
        Marcaje.minutos_tardanza > 0,
    )

    if not admin.get("es_superadmin"):
        query_usuarios = query_usuarios.filter(Usuario.sucursal_id.in_(sucursales_ids))
        query_marcajes = query_marcajes.filter(Marcaje.sucursal_id.in_(sucursales_ids))
        query_justif = query_justif.filter(Marcaje.sucursal_id.in_(sucursales_ids))
        query_entradas = query_entradas.filter(Marcaje.sucursal_id.in_(sucursales_ids))
        query_tardanzas = query_tardanzas.filter(Marcaje.sucursal_id.in_(sucursales_ids))

    total_usuarios = query_usuarios.count()
    marcajes_hoy = query_marcajes.count()
    justif_pendientes = query_justif.count()
    entradas_hoy = query_entradas.count()
    tardanzas_hoy = query_tardanzas.scalar() or 0

    return templates.TemplateResponse(
        request,
        "admin_dashboard.html",
        {
            "admin": admin,
            "total_usuarios": total_usuarios,
            "marcajes_hoy": marcajes_hoy,
            "entradas_hoy": entradas_hoy,
            "justif_pendientes": justif_pendientes,
            "tardanzas_hoy": tardanzas_hoy,
            "hoy": hoy.strftime("%d/%m/%Y"),
        },
    )


# ══════════════════════════════════════════
# GESTIÓN DE SUCURSALES (FASE 2)
# ══════════════════════════════════════════

@router.get("/admin/sucursales", response_class=HTMLResponse)
async def sucursales_page(
    request: Request,
    admin=Depends(requiere_superadmin),
    db: Session = Depends(get_db),
):
    sucursales = db.query(Sucursal).order_by(Sucursal.id).all()
    return templates.TemplateResponse(
        request, "admin_sucursales.html", {"admin": admin, "sucursales": sucursales}
    )


@router.post("/admin/sucursales/crear")
async def crear_sucursal(
    request: Request,
    nombre: str = Form(...),
    codigo: str = Form(...),
    direccion: str = Form(None),
    zona_horaria: str = Form("America/El_Salvador"),
    admin=Depends(requiere_superadmin),
    db: Session = Depends(get_db),
):
    codigo_clean = codigo.strip().upper()
    existe = db.query(Sucursal).filter(Sucursal.codigo == codigo_clean).first()
    if existe:
        raise HTTPException(400, "El código de sucursal ya existe")

    s = Sucursal(
        nombre=nombre.strip(),
        codigo=codigo_clean,
        direccion=direccion.strip() if direccion else None,
        zona_horaria=zona_horaria.strip(),
    )
    db.add(s)
    db.commit()

    _log(
        db,
        admin,
        TipoAccion.CREAR_SUCURSAL if hasattr(TipoAccion, "CREAR_SUCURSAL") else TipoAccion.CREAR_USUARIO,
        {"sucursal_id": s.id, "codigo": codigo_clean, "nombre": s.nombre},
        _ip(request),
    )
    return {"ok": True, "id": s.id}


@router.put("/admin/sucursales/{sucursal_id}/estado")
async def cambiar_estado_sucursal(
    sucursal_id: int,
    request: Request,
    activa: bool = Form(...),
    admin=Depends(requiere_superadmin),
    db: Session = Depends(get_db),
):
    s = db.query(Sucursal).filter(Sucursal.id == sucursal_id).first()
    if not s:
        raise HTTPException(404, "Sucursal no encontrada")

    s.activa = activa
    db.commit()

    accion = (
        TipoAccion.ACTIVAR_SUCURSAL if activa else TipoAccion.DESACTIVAR_SUCURSAL
    ) if hasattr(TipoAccion, "ACTIVAR_SUCURSAL") else TipoAccion.DESACTIVAR_USUARIO

    _log(db, admin, accion, {"sucursal_id": sucursal_id, "nombre": s.nombre}, _ip(request))

    return {"ok": True}


# ══════════════════════════════════════════
# USUARIOS (CRUD)
# ══════════════════════════════════════════

@router.get("/admin/usuarios", response_class=HTMLResponse)
async def lista_usuarios(
    request: Request,
    db: Session = Depends(get_db),
    admin=Depends(obtener_admin_actual),
):
    query = db.query(Usuario)
    if not admin.get("es_superadmin"):
        query = query.filter(Usuario.sucursal_id.in_(admin.get("sucursales_ids", [])))

    usuarios = query.order_by(Usuario.nombre).all()
    sucursales = db.query(Sucursal).filter(Sucursal.activa == True).all()

    return templates.TemplateResponse(
        request,
        "admin_usuarios.html",
        {"admin": admin, "usuarios": usuarios, "sucursales": sucursales},
    )


@router.post("/admin/usuarios/crear")
async def crear_admin_usuario(
    request: Request,
    nombre: str = Form(...),
    apellido: str = Form(...),
    cargo: str = Form(...),
    departamento: str = Form(None),
    email: str = Form(...),
    sucursal_id: int = Form(None),
    turno_entrada: str = Form(None),
    turno_salida: str = Form(None),
    minutos_almuerzo: int = Form(60),
    horas_laborales_dia: int = Form(8),
    db: Session = Depends(get_db),
    admin=Depends(obtener_admin_actual),
):
    if sucursal_id and not admin.get("es_superadmin") and sucursal_id not in admin.get("sucursales_ids", []):
        raise HTTPException(403, "No tiene permisos para asignar usuarios a esta sucursal")

    existe = db.query(Usuario).filter(Usuario.email == email).first()
    if existe:
        raise HTTPException(400, "Email ya registrado")

    t_entrada = datetime.strptime(turno_entrada, "%H:%M").time() if turno_entrada else None
    t_salida = datetime.strptime(turno_salida, "%H:%M").time() if turno_salida else None

    u = Usuario(
        nombre=nombre,
        apellido=apellido,
        cargo=cargo,
        departamento=departamento,
        email=email,
        sucursal_id=sucursal_id,
        turno_entrada=t_entrada,
        turno_salida=t_salida,
        minutos_almuerzo=minutos_almuerzo,
        horas_laborales_dia=horas_laborales_dia,
    )
    db.add(u)
    db.commit()

    _log(
        db,
        admin,
        TipoAccion.CREAR_USUARIO,
        {"usuario_id": u.id, "email": email, "cargo": cargo, "sucursal_id": sucursal_id},
        _ip(request),
    )

    return {"ok": True, "id": u.id}


@router.put("/admin/usuarios/{uid}/horario")
async def actualizar_horario(
    uid: int,
    request: Request,
    turno_entrada: str = Form(None),
    turno_salida: str = Form(None),
    minutos_almuerzo: int = Form(60),
    horas_laborales_dia: int = Form(8),
    db: Session = Depends(get_db),
    admin=Depends(obtener_admin_actual),
):
    u = db.query(Usuario).filter(Usuario.id == uid).first()
    if not u:
        raise HTTPException(404, "Usuario no encontrado")

    if not admin.get("es_superadmin") and u.sucursal_id not in admin.get("sucursales_ids", []):
        raise HTTPException(403, "No tiene permisos para modificar este usuario")

    u.turno_entrada = datetime.strptime(turno_entrada, "%H:%M").time() if turno_entrada else None
    u.turno_salida = datetime.strptime(turno_salida, "%H:%M").time() if turno_salida else None
    u.minutos_almuerzo = minutos_almuerzo
    u.horas_laborales_dia = horas_laborales_dia
    db.commit()

    _log(
        db,
        admin,
        TipoAccion.ACTUALIZAR_HORARIO,
        {
            "usuario_id": uid,
            "turno_entrada": turno_entrada,
            "turno_salida": turno_salida,
        },
        _ip(request),
    )

    return {"ok": True}


@router.put("/admin/usuarios/{uid}/desactivar")
async def desactivar_usuario(
    uid: int,
    request: Request,
    db: Session = Depends(get_db),
    admin=Depends(obtener_admin_actual),
):
    u = db.query(Usuario).filter(Usuario.id == uid).first()
    if not u:
        raise HTTPException(404, "Usuario no encontrado")

    if not admin.get("es_superadmin") and u.sucursal_id not in admin.get("sucursales_ids", []):
        raise HTTPException(403, "No tiene permisos para desactivar este usuario")

    u.activo = False
    db.commit()

    _log(
        db,
        admin,
        TipoAccion.DESACTIVAR_USUARIO,
        {"usuario_id": uid, "email": u.email},
        _ip(request),
    )

    return {"ok": True}


# ══════════════════════════════════════════
# ADMINS (solo superadmin)
# ══════════════════════════════════════════

@router.get("/admin/admins", response_class=HTMLResponse)
@router.get("/admins", response_class=HTMLResponse)
async def admins_page(
    request: Request,
    admin=Depends(requiere_superadmin),
    db: Session = Depends(get_db),
):
    admins = db.query(Administrador).order_by(Administrador.id).all()
    sucursales = db.query(Sucursal).filter(Sucursal.activa == True).all()
    return templates.TemplateResponse(
        request, "admin_admins.html", {"admin": admin, "admins": admins, "sucursales": sucursales}
    )


@router.post("/admin/admins/crear")
@router.post("/admins/crear")
async def admins_crear(
    request: Request,
    nombre: str = Form(...),
    email: str = Form(...),
    password: str = Form(...),
    es_superadmin: str = Form("false"),
    sucursales_ids: list[int] = Form([]),
    admin=Depends(requiere_superadmin),
    db: Session = Depends(get_db),
):
    email_clean = email.strip().lower()
    if len(password) < 8:
        raise HTTPException(400, "La contraseña debe tener al menos 8 caracteres")

    existe = db.query(Administrador).filter(sqlfunc.lower(Administrador.email) == email_clean).first()
    if existe:
        raise HTTPException(409, "Ya existe un administrador con ese email")

    es_super_bool = es_superadmin.lower() == "true"
    a = Administrador(
        nombre=nombre.strip(),
        email=email_clean,
        password_hash=hashear_password(password),
    )
    if hasattr(a, "es_superadmin"):
        a.es_superadmin = es_super_bool

    if sucursales_ids and hasattr(a, "sucursales"):
        sucursales_obj = db.query(Sucursal).filter(Sucursal.id.in_(sucursales_ids)).all()
        a.sucursales = sucursales_obj

    db.add(a)
    db.commit()

    _log(
        db,
        admin,
        TipoAccion.CREAR_ADMINISTRADOR if hasattr(TipoAccion, "CREAR_ADMINISTRADOR") else TipoAccion.CREAR_USUARIO,
        {"email": email_clean, "es_superadmin": es_super_bool, "sucursales_assigned": sucursales_ids},
        _ip(request),
    )
    return {"ok": True}


@router.put("/admin/admins/{admin_id}/password")
@router.put("/admins/{admin_id}/password")
async def admins_password(
    admin_id: int,
    request: Request,
    password: str = Form(...),
    admin=Depends(requiere_superadmin),
    db: Session = Depends(get_db),
):
    if len(password) < 8:
        raise HTTPException(400, "La contraseña debe tener al menos 8 caracteres")

    target_admin = db.query(Administrador).filter(Administrador.id == admin_id).first()
    if not target_admin:
        raise HTTPException(404, "Administrador no encontrado")

    target_admin.password_hash = hashear_password(password)
    db.commit()

    _log(
        db,
        admin,
        TipoAccion.CAMBIAR_PASSWORD_ADMIN if hasattr(TipoAccion, "CAMBIAR_PASSWORD_ADMIN") else TipoAccion.ACTUALIZAR_HORARIO,
        {"target_admin_id": admin_id, "email": target_admin.email},
        _ip(request),
    )
    return {"ok": True}


@router.delete("/admin/admins/{admin_id}")
@router.delete("/admins/{admin_id}")
async def admins_eliminar(
    admin_id: int,
    request: Request,
    admin=Depends(requiere_superadmin),
    db: Session = Depends(get_db),
):
    if admin_id == admin["id"]:
        raise HTTPException(400, "No puedes eliminar tu propia cuenta")

    target_admin = db.query(Administrador).filter(Administrador.id == admin_id).first()
    if not target_admin:
        raise HTTPException(404, "Administrador no encontrado")

    email_borrado = target_admin.email
    db.delete(target_admin)
    db.commit()

    _log(
        db,
        admin,
        TipoAccion.ELIMINAR_ADMIN if hasattr(TipoAccion, "ELIMINAR_ADMIN") else TipoAccion.DESACTIVAR_USUARIO,
        {"deleted_admin_id": admin_id, "email": email_borrado},
        _ip(request),
    )
    return {"ok": True}


# ══════════════════════════════════════════
# LOGS DE AUDITORÍA (solo superadmin)
# ══════════════════════════════════════════

@router.get("/admin/auditoria", response_class=HTMLResponse)
async def auditoria_page(
    request: Request,
    admin=Depends(requiere_superadmin),
):
    return templates.TemplateResponse(
        request, "admin_auditoria.html", {"admin": admin}
    )


@router.get("/admin/auditoria/data")
async def ver_auditoria(
    pagina: int = 1,
    por_pagina: int = 50,
    db: Session = Depends(get_db),
    admin=Depends(requiere_superadmin),
):
    offset = (pagina - 1) * por_pagina
    total = db.query(LogAuditoria).count()
    logs = (
        db.query(LogAuditoria)
        .order_by(LogAuditoria.fecha_hora.desc())
        .offset(offset)
        .limit(por_pagina)
        .all()
    )

    return {
        "total": total,
        "pagina": pagina,
        "logs": [
            {
                "id": l.id,
                "admin_email": l.admin_email,
                "accion": l.accion.value if hasattr(l.accion, "value") else str(l.accion),
                "detalle": l.detalle,
                "ip": l.ip,
                "fecha_hora": l.fecha_hora,
            }
            for l in logs
        ],
    }


# ══════════════════════════════════════════
# HELPERS DE CÁLCULO
# ══════════════════════════════════════════

def calcular_horas_dia(marcajes_dia: list, usuario: Usuario):
    por_tipo = {}
    for m in marcajes_dia:
        por_tipo[m.tipo] = m.fecha_hora

    entrada = por_tipo.get(TipoMarcaje.ENTRADA)
    salida_alm = por_tipo.get(TipoMarcaje.SALIDA_ALMUERZO)
    regreso_alm = por_tipo.get(TipoMarcaje.REGRESO_ALMUERZO)
    salida_final = (
        por_tipo.get(TipoMarcaje.SALIDA)
        or por_tipo.get(TipoMarcaje.SALIDA_IMPREVISTA)
        or por_tipo.get(TipoMarcaje.SALIDA_ENFERMEDAD)
    )

    if not entrada or not salida_final:
        return 0.0, 0.0

    if salida_alm and regreso_alm:
        minutos = (salida_alm - entrada).total_seconds() / 60
        minutos += (salida_final - regreso_alm).total_seconds() / 60
    else:
        total_min = (salida_final - entrada).total_seconds() / 60
        minutos = total_min - (usuario.minutos_almuerzo or 60)

    horas_trabajadas = max(0.0, minutos / 60)
    horas_contratadas = usuario.horas_laborales_dia or 8
    horas_extras = max(0.0, horas_trabajadas - horas_contratadas)
    return round(horas_trabajadas, 2), round(horas_extras, 2)


def agrupar_por_dia(marcajes: list) -> dict:
    dias = {}
    for m in marcajes:
        d = m.fecha_hora.date()
        dias.setdefault(d, []).append(m)
    return dias


# ══════════════════════════════════════════
# JUSTIFICACIONES
# ══════════════════════════════════════════

@router.get("/admin/justificaciones", response_class=HTMLResponse)
async def justificaciones(
    request: Request,
    db: Session = Depends(get_db),
    admin=Depends(obtener_admin_actual),
):
    query = (
        db.query(Marcaje, Usuario)
        .join(Usuario, Marcaje.usuario_id == Usuario.id)
        .filter(
            Marcaje.es_especial == True,
            Marcaje.estado_justificacion == EstadoJustificacion.PENDIENTE,
        )
    )

    if not admin.get("es_superadmin"):
        query = query.filter(Marcaje.sucursal_id.in_(admin.get("sucursales_ids", [])))

    pendientes = query.order_by(Marcaje.fecha_hora.desc()).all()

    return templates.TemplateResponse(
        request,
        "admin_justificaciones.html",
        {"admin": admin, "pendientes": pendientes},
    )


@router.post("/admin/justificaciones/{marcaje_id}/resolver")
async def resolver_justificacion(
    marcaje_id: int,
    request: Request,
    estado: str = Form(...),
    comentario: str = Form(None),
    db: Session = Depends(get_db),
    admin=Depends(obtener_admin_actual),
):
    m = db.query(Marcaje).filter(Marcaje.id == marcaje_id).first()
    if not m:
        raise HTTPException(404, "Marcaje no encontrado")

    if not admin.get("es_superadmin") and m.sucursal_id not in admin.get("sucursales_ids", []):
        raise HTTPException(403, "No tiene permisos para resolver esta justificación")

    m.estado_justificacion = EstadoJustificacion(estado)
    m.comentario_admin = comentario
    db.commit()

    _log(
        db,
        admin,
        TipoAccion.RESOLVER_JUSTIFICACION,
        {
            "marcaje_id": marcaje_id,
            "estado": estado,
            "usuario_id": m.usuario_id,
        },
        _ip(request),
    )

    return {"ok": True}


# ══════════════════════════════════════════
# REPORTES
# ══════════════════════════════════════════

@router.get("/admin/reportes", response_class=HTMLResponse)
async def reportes(
    request: Request,
    db: Session = Depends(get_db),
    admin=Depends(obtener_admin_actual),
):
    query = db.query(Usuario).filter(Usuario.activo == True)
    if not admin.get("es_superadmin"):
        query = query.filter(Usuario.sucursal_id.in_(admin.get("sucursales_ids", [])))

    usuarios = query.order_by(Usuario.nombre).all()
    return templates.TemplateResponse(
        request, "admin_reportes.html", {"admin": admin, "usuarios": usuarios}
    )


def _datos_reporte(
    usuario_id: int, fecha_inicio: date, fecha_fin: date, db: Session
):
    usuario = db.query(Usuario).filter(Usuario.id == usuario_id).first()
    if not usuario:
        raise HTTPException(404, "Usuario no encontrado")

    marcajes = (
        db.query(Marcaje)
        .filter(
            Marcaje.usuario_id == usuario_id,
            sqlfunc.date(Marcaje.fecha_hora) >= fecha_inicio,
            sqlfunc.date(Marcaje.fecha_hora) <= fecha_fin,
        )
        .order_by(Marcaje.fecha_hora)
        .all()
    )

    dias = agrupar_por_dia(marcajes)
    filas = []
    total_trabajadas = 0.0
    total_extras = 0.0
    total_tardanza = 0

    for dia in sorted(dias.keys()):
        horas, extras = calcular_horas_dia(dias[dia], usuario)
        total_trabajadas += horas
        total_extras += extras

        tipos = {m.tipo: m for m in dias[dia]}
        entrada_m = tipos.get(TipoMarcaje.ENTRADA)
        tardanza = entrada_m.minutos_tardanza if entrada_m else 0
        total_tardanza += tardanza

        salida_m = (
            tipos.get(TipoMarcaje.SALIDA)
            or tipos.get(TipoMarcaje.SALIDA_IMPREVISTA)
            or tipos.get(TipoMarcaje.SALIDA_ENFERMEDAD)
        )

        filas.append({
            "fecha": dia.strftime("%d/%m/%Y"),
            "entrada": (
                tipos[TipoMarcaje.ENTRADA].fecha_hora.strftime("%H:%M")
                if TipoMarcaje.ENTRADA in tipos
                else "—"
            ),
            "sal_alm": (
                tipos[TipoMarcaje.SALIDA_ALMUERZO].fecha_hora.strftime("%H:%M")
                if TipoMarcaje.SALIDA_ALMUERZO in tipos
                else "—"
            ),
            "reg_alm": (
                tipos[TipoMarcaje.REGRESO_ALMUERZO].fecha_hora.strftime("%H:%M")
                if TipoMarcaje.REGRESO_ALMUERZO in tipos
                else "—"
            ),
            "salida": salida_m.fecha_hora.strftime("%H:%M") if salida_m else "—",
            "horas": horas,
            "extras": extras,
            "tardanza": tardanza,
        })

    return (
        usuario,
        filas,
        round(total_trabajadas, 2),
        round(total_extras, 2),
        total_tardanza,
    )


@router.get("/admin/reportes/datos")
async def reporte_datos(
    usuario_id: int,
    fecha_inicio: str,
    fecha_fin: str,
    request: Request,
    db: Session = Depends(get_db),
    admin=Depends(obtener_admin_actual),
):
    fi = datetime.strptime(fecha_inicio, "%Y-%m-%d").date()
    ff = datetime.strptime(fecha_fin, "%Y-%m-%d").date()
    usuario, filas, total_h, total_e, total_tardanza = _datos_reporte(
        usuario_id, fi, ff, db
    )

    _log(
        db,
        admin,
        TipoAccion.EXPORTAR_REPORTE,
        {
            "usuario_id": usuario_id,
            "formato": "json",
            "desde": fecha_inicio,
            "hasta": fecha_fin,
        },
        _ip(request),
    )

    return {
        "usuario": f"{usuario.nombre} {usuario.apellido}",
        "cargo": usuario.cargo,
        "filas": filas,
        "total_horas_trabajadas": total_h,
        "total_horas_extras": total_e,
        "total_minutos_tardanza": total_tardanza,
        "horas_contratadas_dia": usuario.horas_laborales_dia or 8,
    }


# ── EXPORTAR PDF ───────────────────────────────────────────────────────

@router.get("/admin/reportes/pdf")
async def exportar_pdf(
    usuario_id: int,
    fecha_inicio: str,
    fecha_fin: str,
    request: Request,
    db: Session = Depends(get_db),
    admin=Depends(obtener_admin_actual),
):
    fi = datetime.strptime(fecha_inicio, "%Y-%m-%d").date()
    ff = datetime.strptime(fecha_fin, "%Y-%m-%d").date()
    usuario, filas, total_h, total_e, total_tardanza = _datos_reporte(
        usuario_id, fi, ff, db
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )
    styles = getSampleStyleSheet()
    elems = []

    elems.append(Paragraph("BIOCORE — Reporte de Marcaciones", styles["Title"]))
    elems.append(Spacer(1, 0.3 * cm))
    elems.append(
        Paragraph(
            f"<b>Empleado:</b> {usuario.nombre} {usuario.apellido} &nbsp;&nbsp; "
            f"<b>Cargo:</b> {usuario.cargo} &nbsp;&nbsp; "
            f"<b>Período:</b> {fi.strftime('%d/%m/%Y')} – {ff.strftime('%d/%m/%Y')}",
            styles["Normal"],
        )
    )
    elems.append(Spacer(1, 0.5 * cm))

    cabecera = [
        "Fecha",
        "Entrada",
        "Sal. Almuerzo",
        "Reg. Almuerzo",
        "Salida",
        "Hrs. Trabajadas",
        "Hrs. Extras",
        "Min. Tardanza",
    ]
    data = [cabecera]
    for f in filas:
        data.append([
            f["fecha"],
            f["entrada"],
            f["sal_alm"],
            f["reg_alm"],
            f["salida"],
            f"{f['horas']}h",
            f"{f['extras']}h",
            f"{f['tardanza']}min" if f["tardanza"] else "—",
        ])
    data.append([
        "TOTAL",
        "",
        "",
        "",
        "",
        f"{total_h}h",
        f"{total_e}h",
        f"{total_tardanza}min",
    ])

    t = Table(data, repeatRows=1)
    t.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            (
                "ROWBACKGROUNDS",
                (0, 1),
                (-1, -2),
                [colors.white, colors.HexColor("#f1f5f9")],
            ),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#00b896")),
            ("TEXTCOLOR", (0, -1), (-1, -1), colors.white),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#e2e8f0")),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("PADDING", (0, 0), (-1, -1), 6),
        ])
    )
    elems.append(t)
    elems.append(Spacer(1, 0.5 * cm))
    elems.append(
        Paragraph(
            f"<b>Total horas trabajadas:</b> {total_h}h &nbsp;&nbsp; "
            f"<b>Total horas extras:</b> {total_e}h &nbsp;&nbsp; "
            f"<b>Total minutos tardanza:</b> {total_tardanza}min",
            styles["Normal"],
        )
    )

    doc.build(elems)
    buf.seek(0)

    _log(
        db,
        admin,
        TipoAccion.EXPORTAR_REPORTE,
        {"usuario_id": usuario_id, "formato": "pdf"},
        _ip(request),
    )

    nombre_archivo = f"reporte_{usuario.apellido}_{fi}_{ff}.pdf"
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={nombre_archivo}"},
    )


# ── EXPORTAR EXCEL ─────────────────────────────────────────────────────

@router.get("/admin/reportes/excel")
async def exportar_excel(
    usuario_id: int,
    fecha_inicio: str,
    fecha_fin: str,
    request: Request,
    db: Session = Depends(get_db),
    admin=Depends(obtener_admin_actual),
):
    fi = datetime.strptime(fecha_inicio, "%Y-%m-%d").date()
    ff = datetime.strptime(fecha_fin, "%Y-%m-%d").date()
    usuario, filas, total_h, total_e, total_tardanza = _datos_reporte(
        usuario_id, fi, ff, db
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Marcaciones"

    verde = PatternFill("solid", fgColor="00B896")
    oscuro = PatternFill("solid", fgColor="0F172A")
    gris = PatternFill("solid", fgColor="F1F5F9")
    blanco = Font(bold=True, color="FFFFFF")
    centro = Alignment(horizontal="center")

    ws.merge_cells("A1:H1")
    ws["A1"] = f"BIOCORE — Reporte: {usuario.nombre} {usuario.apellido} ({fi} / {ff})"
    ws["A1"].font = Font(bold=True, size=13)

    cabecera = [
        "Fecha",
        "Entrada",
        "Sal. Almuerzo",
        "Reg. Almuerzo",
        "Salida",
        "Hrs. Trabajadas",
        "Hrs. Extras",
        "Min. Tardanza",
    ]
    ws.append([])
    ws.append(cabecera)
    fila_cab = ws.max_row
    for col in range(1, 9):
        c = ws.cell(fila_cab, col)
        c.fill = oscuro
        c.font = blanco
        c.alignment = centro

    for i, f in enumerate(filas):
        ws.append([
            f["fecha"],
            f["entrada"],
            f["sal_alm"],
            f["reg_alm"],
            f["salida"],
            f["horas"],
            f["extras"],
            f["tardanza"] if f["tardanza"] else 0,
        ])
        if i % 2 == 1:
            for col in range(1, 9):
                ws.cell(ws.max_row, col).fill = gris

    ws.append(["TOTAL", "", "", "", "", total_h, total_e, total_tardanza])
    fila_tot = ws.max_row
    for col in range(1, 9):
        c = ws.cell(fila_tot, col)
        c.fill = verde
        c.font = blanco
        c.alignment = centro

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 16

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    _log(
        db,
        admin,
        TipoAccion.EXPORTAR_REPORTE,
        {"usuario_id": usuario_id, "formato": "excel"},
        _ip(request),
    )

    nombre = f"reporte_{usuario.apellido}_{fi}_{ff}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={nombre}"},
    )


# ── EXPORTAR CSV ───────────────────────────────────────────────────────

@router.get("/admin/reportes/csv")
async def exportar_csv(
    usuario_id: int,
    fecha_inicio: str,
    fecha_fin: str,
    request: Request,
    db: Session = Depends(get_db),
    admin=Depends(obtener_admin_actual),
):
    fi = datetime.strptime(fecha_inicio, "%Y-%m-%d").date()
    ff = datetime.strptime(fecha_fin, "%Y-%m-%d").date()
    usuario, filas, total_h, total_e, total_tardanza = _datos_reporte(
        usuario_id, fi, ff, db
    )

    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([
        "Fecha",
        "Entrada",
        "Sal. Almuerzo",
        "Reg. Almuerzo",
        "Salida",
        "Hrs. Trabajadas",
        "Hrs. Extras",
        "Min. Tardanza",
    ])
    for f in filas:
        w.writerow([
            f["fecha"],
            f["entrada"],
            f["sal_alm"],
            f["reg_alm"],
            f["salida"],
            f["horas"],
            f["extras"],
            f["tardanza"],
        ])
    w.writerow(["TOTAL", "", "", "", "", total_h, total_e, total_tardanza])

    buf.seek(0)
    _log(
        db,
        admin,
        TipoAccion.EXPORTAR_REPORTE,
        {"usuario_id": usuario_id, "formato": "csv"},
        _ip(request),
    )

    nombre = f"reporte_{usuario.apellido}_{fi}_{ff}.csv"
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={nombre}"},
    )
